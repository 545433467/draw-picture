#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
给 sample.xlsx 随机添加「所属业务」列，然后重新生成 sample.html
用法: python add_business.py
"""
import random
import subprocess
import sys
import os

try:
    import openpyxl
except ImportError:
    sys.exit("请先安装 openpyxl: pip install openpyxl")

XLSX = os.path.join(os.path.dirname(__file__), "sample.xlsx")
OUT  = os.path.join(os.path.dirname(__file__), "sample.html")

# ── 随机业务列表 ──────────────────────────────────────────────────────────────
BUSINESSES = ["用户中心", "订单系统", "支付平台", "数据分析", "内容平台"]

wb = openpyxl.load_workbook(XLSX)
ws = wb.active

# 找现有表头
header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
print("当前列：", header)

# 找或新建「所属业务」列
BIZ_COL_NAME = "所属业务"
if BIZ_COL_NAME in header:
    biz_col = header.index(BIZ_COL_NAME) + 1
    print(f"已存在「所属业务」列（第{biz_col}列），直接覆盖写入")
else:
    biz_col = len(header) + 1
    ws.cell(row=1, column=biz_col, value=BIZ_COL_NAME)
    print(f"新增「所属业务」列（第{biz_col}列）")

# 为每行随机分配业务（约 20% 的行留空，模拟无业务的场景）
rows_written = 0
for row_idx in range(2, ws.max_row + 1):
    name_cell = ws.cell(row=row_idx, column=1).value
    if not name_cell:
        continue
    # 20% 概率不分配业务
    biz = "" if random.random() < 0.20 else random.choice(BUSINESSES)
    ws.cell(row=row_idx, column=biz_col, value=biz)
    rows_written += 1

wb.save(XLSX)
print(f"已为 {rows_written} 行写入业务字段 → {XLSX}")

# ── 重新生成 HTML ─────────────────────────────────────────────────────────────
script = os.path.join(os.path.dirname(__file__), "converter.py")
ret = subprocess.run([sys.executable, script, XLSX, OUT], capture_output=True, text=True)
if ret.stdout:
    print(ret.stdout.strip())
if ret.stderr:
    print("[stderr]", ret.stderr.strip())
if ret.returncode != 0:
    print("converter.py 运行失败，请手动执行：")
    print(f"  python converter.py sample.xlsx sample.html")
else:
    print(f"\n完成！请用浏览器打开 {OUT} 查看效果。")
